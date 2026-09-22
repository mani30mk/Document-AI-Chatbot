"""
RAG Study Assistant — FastAPI Backend (Lightweight Edition)
Uses google-genai SDK directly instead of LangChain to fit within Render 512MB free tier.

Wraps:
  - File upload + multi-format loading (PDF, DOCX, PPTX, TXT)
  - Per-session vector stores (Supabase pgvector primary, in-memory fallback)
  - Multi-turn conversation memory
  - /ask endpoint consumed by the frontend chatbot

Run:
    pip install fastapi uvicorn google-genai python-multipart pypdf
                python-docx python-pptx supabase Pillow
    uvicorn main:app --reload --port 8000
"""

import socket

# Force IPv4 socket resolution to prevent Render container IPv6 handshake/read timeouts to Google APIs
_orig_getaddrinfo = socket.getaddrinfo

def _ipv4_getaddrinfo(host, port, family=0, type=0, proto=0, flags=0):
    if family == 0 or family == socket.AF_UNSPEC:
        family = socket.AF_INET
    return _orig_getaddrinfo(host, port, family, type, proto, flags)

socket.getaddrinfo = _ipv4_getaddrinfo

import os
import uuid
import math
from typing import List

from fastapi import FastAPI, File, UploadFile, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel

# Lightweight Google GenAI SDK (~90 MB vs LangChain's ~420 MB)
from google import genai
from google.genai import types

# Document parsers (direct, no LangChain wrappers)
import docx
from pptx import Presentation
import pypdf

# Supabase client (persistent cloud storage & pgvector)
from supabase.client import Client, create_client

import tempfile
import base64
import io
import gc
import time
import threading
from datetime import datetime, timezone
from PIL import Image
import urllib.request
import urllib.parse
import re
import json

# ── App setup ───────────────────────────────────────────────────────────────
app = FastAPI(title="RAG Study Assistant API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],          # tighten in production
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.api_route("/", methods=["GET", "HEAD"])
def serve_index():
    """Serve the frontend single page app."""
    return FileResponse("index.html")


# ── Global state & Cloud storage setup ─────────────────────────────────────
UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "uploaded_docs")
os.makedirs(UPLOAD_DIR, exist_ok=True)
SESSIONS_FILE = os.path.join(UPLOAD_DIR, "sessions.json")

# In-memory session fallbacks (capped to 1 active vectorstore to conserve Render 512MB RAM)
MAX_CACHED_VECTORSTORES = 1
session_vectorstores: dict[str, dict] = {}  # sid -> {"chunks": [...], "embeddings": [...]}
session_histories: dict[str, list] = {}
session_files: dict[str, list[str]] = {}
session_docs: dict[str, dict[str, str]] = {}
session_slides: dict[str, dict[str, list]] = {}
session_api_keys: dict[str, str] = {}

# Bounded session cache — evict oldest when exceeding limit (re-hydrated from Supabase on demand)
MAX_CACHED_SESSIONS = 5
session_access_order: list[str] = []  # most-recent at end


def touch_session(session_id: str):
    """Mark a session as recently used (move to end of access order)."""
    if session_id in session_access_order:
        session_access_order.remove(session_id)
    session_access_order.append(session_id)


def evict_old_sessions():
    """Evict oldest sessions from in-memory dicts when cache exceeds MAX_CACHED_SESSIONS."""
    while len(session_access_order) > MAX_CACHED_SESSIONS:
        oldest = session_access_order.pop(0)
        session_histories.pop(oldest, None)
        session_files.pop(oldest, None)
        session_docs.pop(oldest, None)
        session_slides.pop(oldest, None)
        session_api_keys.pop(oldest, None)
        session_vectorstores.pop(oldest, None)
        gc.collect()
        print(f"Evicted session {oldest[:8]}... from in-memory cache (will re-hydrate from Supabase on demand).")


# Remote embedding service URL (defaults to deployed microservice)
_env_emb_url = os.getenv("EMBEDDING_SERVICE_URL", "").rstrip("/")
if not _env_emb_url or "document-ai-embeddings" in _env_emb_url:
    EMBEDDING_SERVICE_URL = "https://document-ai-chatbot-7ch2.onrender.com"
else:
    EMBEDDING_SERVICE_URL = _env_emb_url


# Supabase (persistent cloud storage & pgvector)
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY") or os.getenv("SUPABASE_SERVICE_KEY")
supabase_client: Client | None = None

if SUPABASE_URL and SUPABASE_KEY:
    try:
        supabase_client = create_client(SUPABASE_URL, SUPABASE_KEY)
        print("Connected to Supabase for persistent cloud storage & pgvector.")
    except Exception as e:
        print(f"Warning: Could not connect to Supabase: {e}")


# ── Session persistence ────────────────────────────────────────────────────

def load_sessions_from_disk():
    """Dev fallback: Load sessions from local JSON file (not used in production with Supabase)."""
    global session_histories, session_files, session_docs, session_slides
    if os.path.exists(SESSIONS_FILE):
        try:
            with open(SESSIONS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                session_files.update(data.get("files", {}))
                session_docs.update(data.get("docs", {}))
                session_slides.update(data.get("slides", {}))
                for sid, msgs in data.get("histories", {}).items():
                    session_histories[sid] = [
                        {"role": m["role"], "content": m["content"]}
                        for m in msgs
                    ]
            print(f"Loaded {len(session_files)} sessions from disk.")
        except Exception as e:
            print(f"Notice: Could not load sessions from disk ({e})")


def save_sessions_to_disk():
    """Dev fallback: Save sessions to local JSON file (not used in production with Supabase)."""
    try:
        data = {
            "files": session_files,
            "docs": session_docs,
            "slides": session_slides,
            "histories": {
                sid: [
                    {"role": m["role"], "content": m["content"]}
                    for m in msgs
                ]
                for sid, msgs in session_histories.items()
            }
        }
        with open(SESSIONS_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f)
    except Exception as e:
        print(f"Notice: Could not save sessions to disk ({e})")


def load_session(session_id: str) -> dict | None:
    """
    Load a session by session_id from Supabase (production) or local JSON (dev fallback).
    Hydrates in-memory dicts and returns a dict with the session state, or None if not found.
    """
    global session_histories, session_files, session_docs, session_slides

    has_mem = (
        session_id in session_histories
        or session_id in session_files
        or session_id in session_docs
        or session_id in session_slides
    )

    # 1. If Supabase is connected, fetch from chat_sessions table
    if supabase_client:
        try:
            res = (
                supabase_client.table("chat_sessions")
                .select("*")
                .eq("session_id", session_id)
                .execute()
            )
            if res.data and len(res.data) > 0:
                row = res.data[0]
                files = row.get("files") or []
                docs = row.get("docs") or {}
                slides = row.get("slides") or {}
                raw_history = row.get("history") or []

                # Hydrate in-memory request-scoped cache
                session_files[session_id] = files
                session_docs[session_id] = docs
                session_slides[session_id] = slides
                session_histories[session_id] = [
                    {"role": m.get("role", "user"), "content": m.get("content", "")}
                    for m in raw_history
                    if isinstance(m, dict)
                ]
                return {
                    "session_id": session_id,
                    "files": files,
                    "docs": docs,
                    "slides": slides,
                    "history": raw_history,
                }
            elif has_mem:
                raw_history = [
                    {"role": m["role"], "content": m["content"]}
                    for m in session_histories.get(session_id, [])
                ]
                return {
                    "session_id": session_id,
                    "files": session_files.get(session_id, []),
                    "docs": session_docs.get(session_id, {}),
                    "slides": session_slides.get(session_id, {}),
                    "history": raw_history,
                }
            return None
        except Exception as err:
            print(f"Notice: Supabase load_session error for {session_id} ({err}), trying dev fallback...")

    # 2. Dev fallback: check in-memory or load from local JSON disk file
    if not has_mem:
        load_sessions_from_disk()
        has_mem = (
            session_id in session_histories
            or session_id in session_files
            or session_id in session_docs
            or session_id in session_slides
        )

    if has_mem:
        raw_history = [
            {"role": m["role"], "content": m["content"]}
            for m in session_histories.get(session_id, [])
        ]
        return {
            "session_id": session_id,
            "files": session_files.get(session_id, []),
            "docs": session_docs.get(session_id, {}),
            "slides": session_slides.get(session_id, {}),
            "history": raw_history,
        }

    return None


def save_session(
    session_id: str,
    files: list[str] | None = None,
    docs: dict[str, str] | None = None,
    slides: dict[str, list] | None = None,
    history: list | None = None,
):
    """
    Persist session state to Supabase (production) or local JSON (dev fallback).
    """
    global session_histories, session_files, session_docs, session_slides

    # Update in-memory dicts
    if files is not None:
        session_files[session_id] = files
    if docs is not None:
        session_docs[session_id] = docs
    if slides is not None:
        session_slides[session_id] = slides
    if history is not None:
        session_histories[session_id] = history

    cur_files = session_files.get(session_id, [])
    cur_docs = session_docs.get(session_id, {})
    cur_slides = session_slides.get(session_id, {})
    cur_msgs = session_histories.get(session_id, [])

    history_json = [
        {"role": m["role"], "content": m["content"]}
        for m in cur_msgs
    ]

    # 1. Primary: Persist to Supabase chat_sessions table
    if supabase_client:
        try:
            payload = {
                "session_id": session_id,
                "files": cur_files,
                "docs": cur_docs,
                "slides": cur_slides,
                "history": history_json,
                "updated_at": datetime.now(timezone.utc).isoformat(),
            }
            supabase_client.table("chat_sessions").upsert(payload).execute()
            # Evict old sessions after successful Supabase persist
            touch_session(session_id)
            evict_old_sessions()
            return
        except Exception as err:
            print(f"Notice: Supabase save_session error ({err}), falling back to disk...")

    # 2. Dev fallback: Save to local JSON disk file (not used in production)
    touch_session(session_id)
    save_sessions_to_disk()


# ── Shared components ────────────────────────────────────────────────────────
EMBEDDINGS = None
EMBEDDINGS_ERROR = None

RAG_SYSTEM_PROMPT = """You are a helpful study assistant. Answer questions using the document context below when possible.
Be concise, clear, and structure your answer with headings and bullet points where appropriate.
If the answer isn't in the context, say so honestly.

At the very end of your response, on a new line, include a tag with the 2 to 4 word specific academic/technical subject for searching educational lecture videos:
<!-- yt_search: <academic subject keywords> -->

Context:
{context}"""

SUMMARIZE_PROMPT = "You are an expert summarizer. Please provide a comprehensive and concise summary of the following document:\n\n{text}"


# ── Pure Python text splitter (replaces LangChain RecursiveCharacterTextSplitter) ─
def split_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> list[str]:
    """Split text into overlapping chunks, breaking at natural boundaries."""
    if not text or not text.strip():
        return []
    chunks = []
    separators = ["\n\n", "\n", ". ", " "]
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end >= len(text):
            chunks.append(text[start:])
            break
        # Try to break at natural boundary
        best_break = end
        for sep in separators:
            pos = text.rfind(sep, start + chunk_size // 2, end)
            if pos > start:
                best_break = pos + len(sep)
                break
        chunks.append(text[start:best_break])
        start = best_break - chunk_overlap
        if start < 0:
            start = 0
    return [c.strip() for c in chunks if c.strip()]


# ── Gemini API Keys Management & Rotation ──────────────────────────────────
_current_key_index: int = 0


def get_gemini_api_keys() -> list[str]:
    """
    Retrieve Gemini API keys from GEMINI_API_KEY environment variable.
    Supports comma-separated, semicolon-separated, or newline-separated multiple keys.
    """
    raw = os.getenv("GEMINI_API_KEY", "")
    if not raw:
        return []
    keys = [
        k.strip()
        for k in raw.replace("\n", ",").replace(";", ",").split(",")
        if k.strip()
    ]
    return keys


def get_current_api_key() -> str | None:
    """Get the currently active API key."""
    keys = get_gemini_api_keys()
    if not keys:
        return None
    global _current_key_index
    return keys[_current_key_index % len(keys)]


def rotate_api_key() -> str | None:
    """Rotate to the next API key in round-robin fashion. Resets cached embeddings."""
    global _current_key_index, EMBEDDINGS
    keys = get_gemini_api_keys()
    if not keys:
        return None
    _current_key_index = (_current_key_index + 1) % len(keys)
    # Reset cached embeddings so get_embeddings() picks up the new key
    EMBEDDINGS = None
    masked = keys[_current_key_index][:6] + "..." + keys[_current_key_index][-4:] if len(keys[_current_key_index]) > 10 else "***"
    print(f"Rotated to API key #{_current_key_index + 1}/{len(keys)} ({masked})")
    return keys[_current_key_index]


def is_rate_limit_error(exc: Exception) -> bool:
    """Detect if an exception is due to quota exhaustion, high demand, or transient network timeouts."""
    msg = str(exc).lower()
    return any(
        s in msg
        for s in [
            "429",
            "resource_exhausted",
            "resourceexhausted",
            "quota",
            "rate limit",
            "ratelimit",
            "too many requests",
            "503",
            "high demand",
            "overloaded",
            "temporarily unavailable",
            "service unavailable",
            "capacity",
            "timed out",
            "timeout",
            "handshake",
            "ssl",
            "connection error",
            "connection reset",
        ]
    )


AVAILABLE_GEMINI_MODELS: list[str] = []


def get_available_models(api_key: str | None = None) -> list[str]:
    """Query Google API for active models, prioritizing reliable production models first."""
    global AVAILABLE_GEMINI_MODELS
    if AVAILABLE_GEMINI_MODELS:
        return AVAILABLE_GEMINI_MODELS

    candidates = [
        "gemini-flash-latest",
        "gemini-2.5-flash",
        "gemini-flash-lite-latest",
        "gemini-2.0-flash",
    ]

    keys = [api_key] if api_key else get_gemini_api_keys()
    if not keys:
        non_pro = [m for m in candidates if "pro" not in m.lower()]
        pro = [m for m in candidates if "pro" in m.lower()]
        return non_pro + pro

    for key in keys:
        try:
            url = f"https://generativelanguage.googleapis.com/v1beta/models?key={key}"
            req = urllib.request.Request(url, headers={"User-Agent": "DocumentAI/1.0"})
            with urllib.request.urlopen(req, timeout=5.0) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                models = data.get("models", [])
                discovered = []
                for m in models:
                    methods = m.get("supportedGenerationMethods", [])
                    name = m.get("name", "").replace("models/", "")
                    name_lower = name.lower()
                    if (
                        "generateContent" in methods
                        and "gemini" in name_lower
                        and not any(x in name_lower for x in ["tts", "embed", "imagen", "robotics", "computer-use", "thinking", "preview", "exp"])
                    ):
                        discovered.append(name)

                # Prioritize: fast flash models first, then other stable models
                preferred_order = [
                    "gemini-flash-latest",
                    "gemini-2.5-flash",
                    "gemini-flash-lite-latest",
                    "gemini-2.0-flash",
                ]
                sorted_models = []
                for p in preferred_order:
                    if p in discovered and p not in sorted_models:
                        sorted_models.append(p)
                for d in discovered:
                    if d not in sorted_models:
                        sorted_models.append(d)
                for c in candidates:
                    if c not in sorted_models:
                        sorted_models.append(c)

                # Push any "pro" models to the end
                non_pro = [m for m in sorted_models if "pro" not in m.lower()]
                pro = [m for m in sorted_models if "pro" in m.lower()]
                sorted_models = non_pro + pro

                if sorted_models:
                    AVAILABLE_GEMINI_MODELS = sorted_models
                    print(f"Discovered {len(AVAILABLE_GEMINI_MODELS)} available Gemini models: {AVAILABLE_GEMINI_MODELS[:5]}")
                    return AVAILABLE_GEMINI_MODELS
        except Exception as e:
            print(f"Notice: Model discovery via API key ({e}), trying next candidate/defaults.")
            continue

    non_pro = [m for m in candidates if "pro" not in m.lower()]
    pro = [m for m in candidates if "pro" in m.lower()]
    AVAILABLE_GEMINI_MODELS = non_pro + pro
    return AVAILABLE_GEMINI_MODELS


def generate_chat(
    model_name: str,
    api_key: str,
    system_prompt: str,
    history: list[dict],
    user_message: str,
    timeout_s: float = 25.0,
) -> str:
    """
    Call Gemini generate_content with system prompt, chat history, and user message.
    Ensures strict turn alternation (user -> model -> user) to avoid 400 Bad Request.
    """
    client = genai.Client(
        api_key=api_key,
        http_options=types.HttpOptions(timeout=int(timeout_s * 1000)),  # SDK expects milliseconds
    )

    # Build contents list ensuring strict alternation between user and model
    contents = []
    last_role = None
    for msg in history:
        role = "user" if msg.get("role") in ["user", "human"] else "model"
        text = (msg.get("content") or "").strip()
        if not text:
            continue
        if role == last_role:
            contents[-1]["parts"][0]["text"] += "\n\n" + text
        else:
            contents.append({"role": role, "parts": [{"text": text}]})
            last_role = role

    # Append current user question
    if contents and contents[-1]["role"] == "user":
        contents[-1]["parts"][0]["text"] += "\n\n" + user_message
    else:
        contents.append({"role": "user", "parts": [{"text": user_message}]})

    config_kwargs = {
        "system_instruction": system_prompt,
        "temperature": 0.2,
        "automatic_function_calling": types.AutomaticFunctionCallingConfig(disable=True),
    }
    try:
        config_kwargs["thinking_config"] = types.ThinkingConfig(thinking_budget=0)
    except Exception:
        pass

    config = types.GenerateContentConfig(**config_kwargs)

    response = client.models.generate_content(
        model=model_name,
        contents=contents,
        config=config,
    )
    return response.text or ""


def generate_text(model_name: str, api_key: str, prompt: str, timeout_s: float = 25.0) -> str:
    """Simple single-turn text generation with Gemini."""
    client = genai.Client(
        api_key=api_key,
        http_options=types.HttpOptions(timeout=int(timeout_s * 1000)),  # SDK expects milliseconds
    )
    config_kwargs = {
        "temperature": 0.2,
        "automatic_function_calling": types.AutomaticFunctionCallingConfig(disable=True),
    }
    try:
        config_kwargs["thinking_config"] = types.ThinkingConfig(thinking_budget=0)
    except Exception:
        pass

    config = types.GenerateContentConfig(**config_kwargs)
    response = client.models.generate_content(
        model=model_name,
        contents=prompt,
        config=config,
    )
    return response.text or ""


# ── Document loaders (direct, no LangChain) ─────────────────────────────────

def load_docx(path: str) -> list[dict]:
    """Load a DOCX file and return list of document dicts."""
    doc = docx.Document(path)
    text_parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)
    return [{"content": "\n\n".join(text_parts), "metadata": {"source": path}}]


def load_pdf(path: str) -> list[dict]:
    """Load a PDF file using pypdf directly."""
    reader = pypdf.PdfReader(path)
    docs = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append({"content": text, "metadata": {"source": path, "page": i}})
    return docs


def load_txt(path: str) -> list[dict]:
    """Load a plain text file."""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return [{"content": text, "metadata": {"source": path}}]


def extract_shape_images(shape):
    imgs = []
    try:
        if hasattr(shape, "image"):
            imgs.append(shape.image)
        elif getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # Group shape
            for sub_shape in shape.shapes:
                imgs.extend(extract_shape_images(sub_shape))
    except Exception:
        pass
    return imgs


def load_pptx(path: str, session_id: str | None = None) -> tuple[list[dict], list[dict]]:
    # Deduce session_id from path if not explicitly passed
    if not session_id:
        norm_path = os.path.normpath(path)
        parts = norm_path.split(os.sep)
        if "uploaded_docs" in parts:
            idx = parts.index("uploaded_docs")
            if len(parts) > idx + 1:
                session_id = parts[idx + 1]

    prs = Presentation(path)
    docs = []
    slides_data = []
    for slide_idx, slide in enumerate(prs.slides):
        slide_texts = []
        title = ""
        bullets = []
        tables = []
        images = []

        # Check title shape first
        try:
            if slide.shapes.title and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
                slide_texts.append(title)
        except Exception:
            pass

        for shape in slide.shapes:
            try:
                if shape == getattr(slide.shapes, "title", None):
                    continue

                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)
                            if not title and len(text) < 80:
                                title = text
                            elif text != title:
                                bullets.append({
                                    "text": text,
                                    "level": getattr(paragraph, "level", 0),
                                })

                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        if any(row_cells):
                            table_rows.append(row_cells)
                            row_text = " | ".join(c for c in row_cells if c)
                            slide_texts.append(row_text)
                    if table_rows:
                        tables.append(table_rows)

                # Extract pictures/images
                shape_imgs = extract_shape_images(shape)
                for img in shape_imgs:
                    if len(images) >= 2:  # Cap at 2 images per slide (reduce RAM)
                        break
                    try:
                        pil_img = Image.open(io.BytesIO(img.blob))
                        try:
                            if pil_img.width > 500:
                                ratio = 500 / pil_img.width
                                new_size = (500, int(pil_img.height * ratio))
                                pil_img = pil_img.resize(new_size, Image.Resampling.LANCZOS)
                            pil_format = "PNG" if pil_img.mode in ("RGBA", "P") else "JPEG"

                            if session_id:
                                img_dir = os.path.join(UPLOAD_DIR, session_id, "slide_images")
                                os.makedirs(img_dir, exist_ok=True)
                                ext_name = "png" if pil_format == "PNG" else "jpg"
                                img_filename = f"slide_{slide_idx + 1}_img_{len(images) + 1}.{ext_name}"
                                img_file_path = os.path.join(img_dir, img_filename)
                                pil_img.save(img_file_path, format=pil_format, quality=85)
                                images.append(f"/slide_image/{session_id}/{img_filename}")
                            else:
                                buf = io.BytesIO()
                                pil_img.save(buf, format=pil_format, quality=85)
                                b64_str = base64.b64encode(buf.getvalue()).decode("utf-8")
                                mime = f"image/{pil_format.lower()}"
                                images.append(f"data:{mime};base64,{b64_str}")
                        finally:
                            pil_img.close()
                    except Exception as img_err:
                        print(f"Notice: Image extraction ({img_err})")
            except Exception:
                continue

        # Speaker notes
        notes = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass

        if not title and bullets:
            title = bullets.pop(0)
        if not title:
            title = f"Slide {slide_idx + 1}"

        slide_text = "\n".join(slide_texts)
        if notes:
            slide_text += f"\n\nSpeaker Notes: {notes}"

        docs.append({"content": slide_text, "metadata": {"source": path, "slide": slide_idx + 1}})

        slides_data.append({
            "slide_number": slide_idx + 1,
            "title": title if isinstance(title, str) else title.get("text", f"Slide {slide_idx + 1}"),
            "bullets": bullets,
            "tables": tables,
            "images": images,
            "notes": notes,
            "raw_text": slide_text,
        })

    prs = None
    gc.collect()
    return docs, slides_data


def load_file(path: str, ext: str, session_id: str | None = None) -> tuple[list[dict], list[dict] | None]:
    if ext == ".pdf":
        return load_pdf(path), None
    elif ext == ".docx":
        return load_docx(path), None
    elif ext in (".pptx", ".ppt"):
        return load_pptx(path, session_id=session_id)
    elif ext == ".txt":
        return load_txt(path), None
    else:
        raise ValueError(f"Unsupported file extension: {ext}")


# ── Embeddings ───────────────────────────────────────────────────────────────

class RemoteEmbeddings:
    """Embeddings wrapper that calls the remote embedding service via HTTP."""

    def __init__(self, service_url: str):
        self.service_url = service_url.rstrip("/")

    def embed_documents(self, texts: list[str], max_retries: int = 5) -> list[list[float]]:
        """Embed a list of document texts via the remote service with warm-up retry and batching."""
        if not texts:
            return []

        all_embeddings = []
        batch_size = 8
        for i in range(0, len(texts), batch_size):
            batch = texts[i:i + batch_size]
            payload = json.dumps({"texts": batch}).encode("utf-8")

            last_err = None
            for attempt in range(max_retries):
                try:
                    req = urllib.request.Request(
                        f"{self.service_url}/embed",
                        data=payload,
                        headers={"Content-Type": "application/json"},
                        method="POST",
                    )
                    with urllib.request.urlopen(req, timeout=60) as resp:
                        data = json.loads(resp.read().decode("utf-8"))
                    all_embeddings.extend(data["embeddings"])
                    last_err = None
                    break
                except Exception as e:
                    last_err = e
                    print(f"Remote embedding batch {i//batch_size + 1} attempt {attempt + 1}/{max_retries} failed ({e}). Waking up / retrying in 5s...")
                    time.sleep(5)

            if last_err is not None:
                raise last_err

        return all_embeddings

    def embed_query(self, text: str, max_retries: int = 4) -> list[float]:
        """Embed a single query text via the remote service with warm-up retry."""
        payload = json.dumps({"texts": [text]}).encode("utf-8")
        last_err = None
        for attempt in range(max_retries):
            try:
                req = urllib.request.Request(
                    f"{self.service_url}/embed",
                    data=payload,
                    headers={"Content-Type": "application/json"},
                    method="POST",
                )
                with urllib.request.urlopen(req, timeout=30) as resp:
                    data = json.loads(resp.read().decode("utf-8"))
                return data["embeddings"][0]
            except Exception as e:
                last_err = e
                print(f"Remote query embedding attempt {attempt + 1}/{max_retries} failed ({e}). Retrying in 4s...")
                time.sleep(4)
        raise last_err


class GeminiEmbeddings:
    """Lightweight embeddings using google-genai SDK directly (no LangChain)."""

    def __init__(self, api_key: str, model: str = "gemini-embedding-001"):
        self.client = genai.Client(api_key=api_key)
        self.model = model

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        """Embed a list of texts using Gemini embedding API with candidate model fallback."""
        candidate_models = [self.model, "gemini-embedding-001", "gemini-embedding-2"]
        unique_models = []
        for m in candidate_models:
            if m and m not in unique_models:
                unique_models.append(m)

        last_err = None
        for m in unique_models:
            try:
                all_embeddings = []
                for i in range(0, len(texts), 100):
                    batch = texts[i:i + 100]
                    result = self.client.models.embed_content(
                        model=m,
                        contents=batch,
                    )
                    all_embeddings.extend([list(e.values) for e in result.embeddings])
                self.model = m
                return all_embeddings
            except Exception as e:
                last_err = e
                print(f"Notice: Model {m} failed for embed_content ({e}), trying next candidate...")
                continue
        raise last_err

    def embed_query(self, text: str) -> list[float]:
        """Embed a single query text."""
        return self.embed_documents([text])[0]


class UnifiedEmbeddings:
    """
    Dedicated FastEmbed microservice (384-dim, 0 MB local RAM, matches Supabase vector(384)).
    """

    def __init__(self, remote_url: str | None = None, api_key: str | None = None):
        self.primary_url = (remote_url or "https://document-ai-chatbot-7ch2.onrender.com").rstrip("/")
        self.verified_url = "https://document-ai-chatbot-7ch2.onrender.com"
        self._remote = RemoteEmbeddings(self.primary_url)
        self._gemini = GeminiEmbeddings(api_key) if api_key else None

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        try:
            return self._remote.embed_documents(texts)
        except Exception as e:
            print(f"Embedding service ({self.primary_url}) error: {e}")

        if self.primary_url != self.verified_url:
            try:
                print(f"Retrying with verified 2nd microservice: {self.verified_url}")
                return RemoteEmbeddings(self.verified_url).embed_documents(texts)
            except Exception as e:
                print(f"Verified 2nd microservice error: {e}")

        if self._gemini:
            try:
                return self._gemini.embed_documents(texts)
            except Exception as e:
                print(f"Gemini fallback embed error: {e}")

        raise ValueError("Could not embed documents. Please verify https://document-ai-chatbot-7ch2.onrender.com is reachable.")

    def embed_query(self, text: str) -> list[float]:
        try:
            return self._remote.embed_query(text)
        except Exception as e:
            print(f"Embedding service query error ({self.primary_url}): {e}")

        if self.primary_url != self.verified_url:
            try:
                return RemoteEmbeddings(self.verified_url).embed_query(text)
            except Exception as e:
                print(f"Verified 2nd microservice query error: {e}")

        if self._gemini:
            try:
                return self._gemini.embed_query(text)
            except Exception as e:
                print(f"Gemini fallback query error: {e}")

        raise ValueError("Could not embed query. Please verify https://document-ai-chatbot-7ch2.onrender.com is reachable.")


def get_embeddings():
    """
    Returns an embeddings instance with fallback priority:
      1. Remote embedding service (EMBEDDING_SERVICE_URL) — 0 MB local RAM, saves tokens, 384-dim
      2. Google cloud embeddings (google-genai SDK) — 0 MB extra RAM, uses Gemini quota
    """
    global EMBEDDINGS
    if EMBEDDINGS is not None:
        return EMBEDDINGS

    remote_url = EMBEDDING_SERVICE_URL or "https://document-ai-chatbot-7ch2.onrender.com"
    key = get_current_api_key()

    EMBEDDINGS = UnifiedEmbeddings(remote_url=remote_url, api_key=key)
    return EMBEDDINGS


# ── Simple in-memory vector search (replaces ChromaDB) ──────────────────────

def cosine_similarity(a: list[float], b: list[float]) -> float:
    """Compute cosine similarity between two vectors."""
    dot = sum(x * y for x, y in zip(a, b))
    norm_a = math.sqrt(sum(x * x for x in a))
    norm_b = math.sqrt(sum(x * x for x in b))
    if norm_a == 0 or norm_b == 0:
        return 0.0
    return dot / (norm_a * norm_b)


def search_vectors(query_embedding: list[float], store: dict, k: int = 5) -> list[dict]:
    """Search in-memory vector store for top-k most similar documents."""
    if not store or "embeddings" not in store or not store["embeddings"]:
        return []
    scored = []
    for i, emb in enumerate(store["embeddings"]):
        sim = cosine_similarity(query_embedding, emb)
        scored.append((sim, store["chunks"][i]))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [item[1] for item in scored[:k]]


def search_supabase_vectors(query_embedding: list[float], session_id: str, k: int = 5) -> list[dict]:
    """Search Supabase pgvector using the match_documents RPC function."""
    if not supabase_client:
        return []
    try:
        result = supabase_client.rpc("match_documents", {
            "query_embedding": query_embedding,
            "match_count": k,
            "filter": {"session_id": session_id},
        }).execute()
        if result.data:
            return [{"content": r.get("content", ""), "metadata": r.get("metadata", {})} for r in result.data]
    except Exception as err:
        print(f"Supabase vector search error ({err})")
    return []


# ── Request / Response models ────────────────────────────────────────────────
class AskRequest(BaseModel):
    session_id: str
    question: str


class SummarizeRequest(BaseModel):
    session_id: str
    filename: str
    text: str | None = None


class YouTubeVideo(BaseModel):
    id: str
    title: str
    channel: str
    duration: str
    url: str
    thumbnail: str


class AskResponse(BaseModel):
    answer: str
    session_id: str
    sources_used: int
    youtube_sources: List[YouTubeVideo] = []


class SessionResponse(BaseModel):
    session_id: str
    files: List[str]


def search_youtube(query: str, max_results: int = 3) -> list[dict]:
    """Search YouTube for educational tutorials matching the query without API key."""
    try:
        clean_query = query.strip()
        search_terms = clean_query + " tutorial lecture"
        url = "https://www.youtube.com/results?search_query=" + urllib.parse.quote(search_terms)
        req = urllib.request.Request(
            url,
            headers={
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/115.0.0.0 Safari/537.36"
            },
        )
        with urllib.request.urlopen(req, timeout=3.5) as response:
            html = response.read().decode("utf-8")
            match = re.search(r"var ytInitialData = ({.*?});</script>", html)
            if not match:
                return []
            data = json.loads(match.group(1))
            contents = data["contents"]["twoColumnSearchResultsRenderer"]["primaryContents"]["sectionListRenderer"]["contents"]
            videos = []
            for section in contents:
                items = section.get("itemSectionRenderer", {}).get("contents", [])
                for item in items:
                    v = item.get("videoRenderer")
                    if v and "videoId" in v:
                        vid_id = v.get("videoId")
                        title = v.get("title", {}).get("runs", [{}])[0].get("text", "")
                        channel = v.get("ownerText", {}).get("runs", [{}])[0].get("text", "")
                        duration = v.get("lengthText", {}).get("simpleText", "")
                        if vid_id and title and "#shorts" not in title.lower() and "#short" not in title.lower():
                            videos.append({
                                "id": vid_id,
                                "title": title,
                                "channel": channel,
                                "duration": duration,
                                "url": f"https://www.youtube.com/watch?v={vid_id}",
                                "thumbnail": f"https://img.youtube.com/vi/{vid_id}/mqdefault.jpg",
                            })
                            if len(videos) >= max_results:
                                break
                if len(videos) >= max_results:
                    break
            return videos
    except Exception as e:
        print(f"YouTube search notice: {e}")
        return []


# ── Endpoints ────────────────────────────────────────────────────────────────

@app.post("/session/new", response_model=SessionResponse)
def new_session():
    """Create a new chat/upload session."""
    sid = str(uuid.uuid4())
    save_session(sid, files=[], docs={}, slides={}, history=[])
    return SessionResponse(session_id=sid, files=[])


@app.get("/session/{session_id}")
def get_session_state(session_id: str):
    """Retrieve full session state (files, slides, docs, chat history) for page reload."""
    sess = load_session(session_id)
    if not sess:
        raise HTTPException(status_code=404, detail="Session not found.")

    history_list = sess.get("history", [])

    return {
        "session_id": session_id,
        "files": sess.get("files", []),
        "slides": sess.get("slides", {}),
        "docs": sess.get("docs", {}),
        "history": history_list,
        "history_turns": len(history_list) // 2,
    }


@app.on_event("startup")
def startup_prewarm():
    """Load sessions on startup."""
    load_sessions_from_disk()
    embed_mode = "remote service" if EMBEDDING_SERVICE_URL else ("cloud API" if get_gemini_api_keys() else "none")
    print(f"Startup complete. Embeddings: {embed_mode}.")


@app.get("/raw/{session_id}/{filename}")
def get_raw_file(session_id: str, filename: str):
    """Serve original uploaded file directly for Office Viewer or download."""
    file_path = os.path.join(UPLOAD_DIR, session_id, filename)
    if os.path.isfile(file_path):
        media_type = None
        lower_name = filename.lower()
        if lower_name.endswith((".pptx", ".ppt")):
            media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        elif lower_name.endswith(".pdf"):
            media_type = "application/pdf"
        elif lower_name.endswith((".docx", ".doc")):
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif lower_name.endswith((".xlsx", ".xls")):
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        elif lower_name.endswith((".txt", ".csv", ".md", ".py", ".json", ".log")):
            media_type = "text/plain; charset=utf-8"
        return FileResponse(
            file_path,
            filename=filename,
            media_type=media_type,
            content_disposition_type="inline",
            headers={
                "Access-Control-Allow-Origin": "*",
                "Cache-Control": "public, max-age=3600",
            },
        )

    if supabase_client:
        try:
            from fastapi.responses import RedirectResponse
            pub_url = supabase_client.storage.from_("documents").get_public_url(f"{session_id}/{filename}")
            if pub_url:
                return RedirectResponse(pub_url)
        except Exception:
            pass
    raise HTTPException(status_code=404, detail="File not found.")


@app.get("/slide_image/{session_id}/{filename}")
def get_slide_image(session_id: str, filename: str):
    """Serve slide images extracted from PPTX presentations."""
    clean_session_id = os.path.basename(session_id)
    clean_filename = os.path.basename(filename)

    expected_dir = os.path.abspath(os.path.join(UPLOAD_DIR, clean_session_id, "slide_images"))
    img_path = os.path.abspath(os.path.join(expected_dir, clean_filename))

    if not img_path.startswith(expected_dir) or not os.path.isfile(img_path):
        raise HTTPException(status_code=404, detail="Slide image not found.")

    ext = os.path.splitext(clean_filename)[1].lower()
    media_type = "image/png" if ext == ".png" else "image/jpeg"
    return FileResponse(
        img_path,
        media_type=media_type,
        headers={
            "Access-Control-Allow-Origin": "*",
            "Cache-Control": "public, max-age=86400",
        },
    )


@app.post("/upload/{session_id}")
async def upload_files(
    session_id: str,
    files: List[UploadFile] = File(...),
):
    """
    Upload one or more study files into the session's vector store.
    Supports PDF, DOCX, PPTX, TXT.
    """
    load_session(session_id)
    if session_id not in session_histories:
        session_histories[session_id] = []
        session_files[session_id] = []
        session_docs[session_id] = {}
        session_slides[session_id] = {}

    all_docs = []
    uploaded_names = []
    session_dir = os.path.join(UPLOAD_DIR, session_id)
    os.makedirs(session_dir, exist_ok=True)

    ALLOWED_EXTENSIONS = {".pdf", ".pptx", ".ppt", ".docx", ".doc", ".txt"}
    for uf in files:
        ext = os.path.splitext(uf.filename)[1].lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file type '{ext}' for file '{uf.filename}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
            )
        file_bytes = await uf.read()

        saved_path = os.path.join(session_dir, uf.filename)
        with open(saved_path, "wb") as f:
            f.write(file_bytes)

        try:
            docs, slides_info = load_file(saved_path, ext, session_id=session_id)
            all_docs.extend(docs)
            uploaded_names.append(uf.filename)
            if slides_info:
                session_slides.setdefault(session_id, {})[uf.filename] = slides_info

            # Store full text for summarization
            full_text = "\n\n".join(d["content"] for d in docs)
            session_docs.setdefault(session_id, {})[uf.filename] = full_text

            # If Supabase is connected, store the raw file in Supabase Storage
            if supabase_client:
                try:
                    supabase_client.storage.from_("documents").upload(
                        path=f"{session_id}/{uf.filename}",
                        file=file_bytes,
                        file_options={"upsert": "true"},
                    )
                except Exception as s_err:
                    print(f"Supabase storage upload note: {s_err}")
        except Exception as e:
            return {"error": f"Failed to parse {uf.filename}: {str(e)}", "skipped": uf.filename}

    if not all_docs:
        raise HTTPException(
            status_code=400,
            detail="No readable text could be extracted from the uploaded document(s). Please verify the file contains readable text.",
        )

    # Split all documents into chunks
    all_texts = []
    all_metadata = []
    for d in all_docs:
        chunks = split_text(d["content"])
        for chunk in chunks:
            all_texts.append(chunk)
            all_metadata.append({**d.get("metadata", {}), "session_id": session_id})

    # Embedding & Indexing with Key Rotation on Rate Limit (429)
    keys = get_gemini_api_keys()
    max_attempts = max(len(keys), 1) if keys else 1
    stored_in_supabase = False
    indexing_success = False
    last_emb_error = None

    for attempt in range(max_attempts):
        try:
            curr_embeddings = get_embeddings()
            vectors = curr_embeddings.embed_documents(all_texts)

            # Persist in Supabase pgvector if configured
            if supabase_client:
                try:
                    # Insert documents with embeddings into Supabase
                    rows = []
                    for i, (text, meta, vec) in enumerate(zip(all_texts, all_metadata, vectors)):
                        rows.append({
                            "content": text,
                            "metadata": meta,
                            "embedding": vec,
                        })
                    # Batch insert
                    supabase_client.table("documents").insert(rows).execute()
                    stored_in_supabase = True
                except Exception as err:
                    if is_rate_limit_error(err):
                        raise  # Let outer retry loop handle 429
                    print(f"Warning: Supabase vector insert failed ({err}), falling back to in-memory.")
                    stored_in_supabase = False

            if not stored_in_supabase:
                # In-memory fallback vector store
                if len(session_vectorstores) >= MAX_CACHED_VECTORSTORES:
                    oldest_sid = next(iter(session_vectorstores))
                    del session_vectorstores[oldest_sid]
                if session_id in session_vectorstores:
                    session_vectorstores[session_id]["chunks"].extend(
                        [{"content": t, "metadata": m} for t, m in zip(all_texts, all_metadata)]
                    )
                    session_vectorstores[session_id]["embeddings"].extend(vectors)
                else:
                    session_vectorstores[session_id] = {
                        "chunks": [{"content": t, "metadata": m} for t, m in zip(all_texts, all_metadata)],
                        "embeddings": vectors,
                    }

            indexing_success = True
            break
        except Exception as emb_err:
            last_emb_error = emb_err
            if is_rate_limit_error(emb_err) and keys:
                print(f"Embedding rate limit / 429 with current key ({emb_err}). Rotating key...")
                rotate_api_key()
                continue
            else:
                print(f"Embedding / indexing error ({emb_err})")
                raise HTTPException(
                    status_code=500,
                    detail=f"Failed to embed/index documents: {str(emb_err)}",
                )

    if not indexing_success:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to embed/index after {max_attempts} key rotations: {str(last_emb_error)}",
        )

    # Clean up upload buffers and trigger garbage collection immediately
    chunks_count = len(all_texts)
    del all_docs
    del all_texts
    del vectors
    gc.collect()

    session_files[session_id].extend(uploaded_names)
    save_session(session_id)

    return {
        "message": f"Uploaded {len(uploaded_names)} file(s), indexed {chunks_count} chunks.",
        "files": session_files[session_id],
        "chunks": chunks_count,
        "slides": session_slides.get(session_id, {}),
        "docs": session_docs.get(session_id, {}),
        "storage": "supabase" if stored_in_supabase else "in-memory",
    }


@app.get("/slides/{session_id}/{filename}")
def get_slides(session_id: str, filename: str):
    """Retrieve structured slides for PPT viewer."""
    if session_id not in session_slides:
        load_session(session_id)
    if session_id not in session_slides or filename not in session_slides[session_id]:
        raise HTTPException(status_code=404, detail="Slides not found.")
    return {
        "filename": filename,
        "slides": session_slides[session_id][filename],
    }


@app.get("/document/{session_id}/{filename}")
def get_document_text(session_id: str, filename: str):
    """Retrieve extracted document text for DOCX/TXT viewer."""
    if session_id not in session_docs:
        load_session(session_id)
    if session_id not in session_docs or filename not in session_docs[session_id]:
        raise HTTPException(status_code=404, detail="Document text not found.")
    return {
        "filename": filename,
        "text": session_docs[session_id][filename],
    }


@app.post("/ask", response_model=AskResponse)
async def ask(req: AskRequest):
    """
    Ask a question. Uses RAG + multi-turn conversation history.
    """
    sid = req.session_id

    if sid not in session_histories:
        load_session(sid)
    if sid not in session_histories:
        raise HTTPException(status_code=404, detail="Session not found.")

    # Retrieve relevant context via vector search
    retrieved_docs = []
    try:
        embeddings = get_embeddings()
        query_vec = embeddings.embed_query(req.question)

        # Try Supabase pgvector first
        if supabase_client:
            retrieved_docs = search_supabase_vectors(query_vec, sid, k=5)

        # Fallback to in-memory vector store
        if not retrieved_docs and sid in session_vectorstores:
            retrieved_docs = search_vectors(query_vec, session_vectorstores[sid], k=5)
    except Exception as r_err:
        print(f"Vector search error ({r_err}), using raw document text fallback.")

    context = "\n\n".join(d["content"] for d in retrieved_docs) if retrieved_docs else ""
    if not context and sid in session_docs:
        context = "\n\n".join(list(session_docs[sid].values()))[:4000]

    models_to_try = get_available_models()[:5]
    keys = get_gemini_api_keys()
    if not keys:
        raise HTTPException(
            status_code=500,
            detail="GEMINI_API_KEY is not configured. Please set GEMINI_API_KEY in Render or in your .env file.",
        )

    # Build system prompt with context
    system_prompt = RAG_SYSTEM_PROMPT.format(context=context)
    history = session_histories[sid]

    answer = None
    last_error = None
    start_time = time.monotonic()
    MAX_BUDGET_SECONDS = 35.0

    for m_name in models_to_try:
        remaining = MAX_BUDGET_SECONDS - (time.monotonic() - start_time)
        if remaining <= 2:
            print(f"Time budget exceeded ({MAX_BUDGET_SECONDS - remaining:.1f}s), breaking model loop in /ask...")
            break
        for _ in range(len(keys)):
            remaining = MAX_BUDGET_SECONDS - (time.monotonic() - start_time)
            if remaining <= 2:
                print(f"Time budget exceeded ({MAX_BUDGET_SECONDS - remaining:.1f}s), breaking key loop in /ask...")
                break
            current_key = get_current_api_key()
            try:
                answer = generate_chat(m_name, current_key, system_prompt, history, req.question, timeout_s=min(remaining, 15.0))
                if answer:
                    break
            except Exception as llm_err:
                last_error = llm_err
                if is_rate_limit_error(llm_err):
                    print(f"Transient error (quota/high demand) on model {m_name}: {llm_err}. Trying next...")
                    if len(keys) > 1:
                        rotate_api_key()
                        continue
                    else:
                        break  # Immediately try next candidate model
                else:
                    print(f"Model {m_name} failed ({llm_err}), falling back to next model candidate...")
                    break
        if answer:
            break

    if not answer:
        err_msg = f"AI model error: {str(last_error)}"
        if time.monotonic() - start_time >= MAX_BUDGET_SECONDS:
            err_msg = f"AI request timed out after {MAX_BUDGET_SECONDS:.0f}s: {str(last_error)}"
        raise HTTPException(status_code=500, detail=err_msg)

    yt_results = []
    try:
        yt_query = req.question
        if answer:
            yt_match = re.search(r"<!--\s*yt_search:\s*(.*?)\s*-->", answer, re.IGNORECASE)
            if yt_match:
                extracted = yt_match.group(1).strip()
                if extracted:
                    yt_query = extracted
                answer = re.sub(r"<!--\s*yt_search:\s*.*?\s*-->", "", answer).strip()
            else:
                clean_q = re.sub(r"^(what is|what are|explain|how to|describe|define)\s+", "", req.question, flags=re.I).strip(" ?.")
                doc_title = ""
                if sid in session_files and session_files[sid]:
                    doc_title = os.path.splitext(session_files[sid][0])[0].replace("_", " ").replace("-", " ")
                    doc_title = re.sub(r"\b(unit|chapter|notes|doc|file|pdf|part|presentation)\s*\d*\b", "", doc_title, flags=re.I).strip()
                if doc_title and len(clean_q) < 30:
                    yt_query = f"{doc_title} {clean_q}"
                else:
                    yt_query = clean_q

        yt_results = search_youtube(yt_query, max_results=3)
    except Exception as yt_err:
        print(f"YouTube search notice: {yt_err}")

    # Persist turn to history
    session_histories[sid].append({"role": "user", "content": req.question})
    session_histories[sid].append({"role": "bot", "content": answer})

    # Keep last 20 messages (10 turns) to avoid token bloat
    if len(session_histories[sid]) > 20:
        session_histories[sid] = session_histories[sid][-20:]

    save_session(sid)

    return AskResponse(
        answer=answer,
        session_id=sid,
        sources_used=len(retrieved_docs),
        youtube_sources=[YouTubeVideo(**v) for v in yt_results],
    )


def normalize_doc_name(name: str) -> str:
    """Normalize filename for fuzzy matching (case, hyphens, underscores, spaces)."""
    return re.sub(r"[\s\-_]+", "", name.lower())


@app.post("/summarize")
async def summarize(req: SummarizeRequest):
    """Summarize a specific uploaded file with multi-tier fallback."""
    sid = req.session_id
    fname = req.filename
    text = (req.text or "").strip()

    # Ensure session state is loaded from Supabase if not in memory
    if sid not in session_docs and sid not in session_slides:
        load_session(sid)

    # Tier 1: Client provided extracted text directly
    if not text:
        # Tier 2: Exact match in session_docs
        if sid in session_docs and fname in session_docs[sid]:
            text = session_docs[sid][fname]

    if not text:
        # Tier 3: Normalized filename match in session_docs[sid]
        target_norm = normalize_doc_name(fname)
        if sid in session_docs:
            for k, doc_text in session_docs[sid].items():
                if normalize_doc_name(k) == target_norm:
                    text = doc_text
                    break

    if not text:
        # Tier 4: Search across all sessions in session_docs
        target_norm = normalize_doc_name(fname)
        for s_id, doc_dict in session_docs.items():
            for k, doc_text in doc_dict.items():
                if normalize_doc_name(k) == target_norm:
                    text = doc_text
                    break
            if text:
                break

    if not text:
        # Tier 5: Check disk in UPLOAD_DIR for the file
        target_norm = normalize_doc_name(fname)
        found_file = None
        # Check session dir first
        session_dir = os.path.join(UPLOAD_DIR, sid)
        if os.path.isdir(session_dir):
            for disk_file in os.listdir(session_dir):
                if disk_file == fname or normalize_doc_name(disk_file) == target_norm:
                    found_file = os.path.join(session_dir, disk_file)
                    break

        # Check entire UPLOAD_DIR recursively if not found in session dir
        if not found_file and os.path.isdir(UPLOAD_DIR):
            for root, _, disk_files in os.walk(UPLOAD_DIR):
                for disk_file in disk_files:
                    if disk_file == fname or normalize_doc_name(disk_file) == target_norm:
                        found_file = os.path.join(root, disk_file)
                        break
                if found_file:
                    break

        if found_file:
            try:
                ext = os.path.splitext(found_file)[1].lower()
                docs, _ = load_file(found_file, ext, session_id=sid)
                if docs:
                    text = "\n\n".join(d["content"] for d in docs)
                    session_docs.setdefault(sid, {})[fname] = text
                    save_session(sid)
            except Exception as load_err:
                print(f"Notice: Could not load file from disk for summarize ({load_err})")

    if not text:
        # Tier 6: Single document fallback in session
        if sid in session_docs and len(session_docs[sid]) == 1:
            text = list(session_docs[sid].values())[0]

    if not text:
        # Tier 7: Check slides if PPT
        if sid in session_slides:
            slides_data = session_slides[sid].get(fname)
            if not slides_data:
                target_norm = normalize_doc_name(fname)
                for k, s_list in session_slides[sid].items():
                    if normalize_doc_name(k) == target_norm:
                        slides_data = s_list
                        break
            if slides_data:
                slide_parts = []
                for s in slides_data:
                    title = s.get("title", "")
                    raw = s.get("raw_text", "")
                    bullets = s.get("bullets", [])
                    bullet_text = "\n".join(
                        b if isinstance(b, str) else b.get("text", "") for b in bullets
                    )
                    part = f"Slide {s.get('slide_number', '')}: {title}\n{raw}\n{bullet_text}".strip()
                    if part:
                        slide_parts.append(part)
                if slide_parts:
                    text = "\n\n".join(slide_parts)

    if not text:
        raise HTTPException(
            status_code=404,
            detail=f"File '{fname}' not found in session and no text could be extracted.",
        )

    prompt = SUMMARIZE_PROMPT.format(text=text[:20000])

    models_to_try = get_available_models()[:5]
    keys = get_gemini_api_keys()
    if not keys:
        raise HTTPException(
            status_code=500,
            detail="GEMINI_API_KEY is not configured. Please set GEMINI_API_KEY in Render or in your .env file.",
        )

    summary = None
    last_error = None
    start_time = time.monotonic()
    MAX_BUDGET_SECONDS = 35.0

    for m_name in models_to_try:
        remaining = MAX_BUDGET_SECONDS - (time.monotonic() - start_time)
        if remaining <= 2:
            print(f"Time budget exceeded ({MAX_BUDGET_SECONDS - remaining:.1f}s), breaking model loop in /summarize...")
            break
        for _ in range(len(keys)):
            remaining = MAX_BUDGET_SECONDS - (time.monotonic() - start_time)
            if remaining <= 2:
                print(f"Time budget exceeded ({MAX_BUDGET_SECONDS - remaining:.1f}s), breaking key loop in /summarize...")
                break
            current_key = get_current_api_key()
            try:
                summary = generate_text(m_name, current_key, prompt, timeout_s=min(remaining, 15.0))
                if summary:
                    break
            except Exception as e:
                last_error = e
                if is_rate_limit_error(e):
                    print(f"Summarize transient error (quota/high demand) on model {m_name}: {e}. Trying next...")
                    if len(keys) > 1:
                        rotate_api_key()
                        continue
                    else:
                        break  # Try next model candidate
                else:
                    print(f"Summarize model {m_name} failed ({e}), falling back to next model candidate...")
                    break
        if summary:
            break

    if not summary:
        err_msg = f"AI model error: {str(last_error)}"
        if time.monotonic() - start_time >= MAX_BUDGET_SECONDS:
            err_msg = f"AI summarization timed out after {MAX_BUDGET_SECONDS:.0f}s: {str(last_error)}"
        raise HTTPException(status_code=500, detail=err_msg)

    return {"summary": summary}



@app.delete("/session/{session_id}/history")
def clear_history(session_id: str):
    """Clear conversation history for a session (keeps files/vectorstore)."""
    if session_id not in session_histories:
        load_session(session_id)
    if session_id not in session_histories:
        raise HTTPException(status_code=404, detail="Session not found.")
    session_histories[session_id] = []
    save_session(session_id)
    return {"message": "Conversation history cleared."}


@app.api_route("/health", methods=["GET", "HEAD"])
def health():
    return {"status": "ok"}
