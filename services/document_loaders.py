"""
Multi-format document parsing and text chunking (PDF, DOCX, PPTX, TXT).
Direct parser implementation without LangChain wrappers to minimize memory usage.
"""

import base64
import gc
import io
import os
from typing import Optional, Tuple, List, Dict

import docx
from pptx import Presentation
import pypdf
from PIL import Image

from config import UPLOAD_DIR


def split_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> list[str]:
    """Split text into overlapping chunks, breaking at natural boundaries."""
    if not text or not text.strip():
        return []
    chunks = []
    separators = ["\n\n", "\n", ". ", " "]
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end >= len(text):
            chunks.append(text[start:])
            break
        # Try to break at natural boundary
        best_break = end
        for sep in separators:
            pos = text.rfind(sep, start + chunk_size // 2, end)
            if pos > start:
                best_break = pos + len(sep)
                break
        chunks.append(text[start:best_break])
        start = best_break - chunk_overlap
        if start < 0:
            start = 0
    return [c.strip() for c in chunks if c.strip()]


def load_docx(path: str) -> list[dict]:
    """Load a DOCX file and return list of document dicts."""
    doc = docx.Document(path)
    text_parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)
    return [{"content": "\n\n".join(text_parts), "metadata": {"source": path}}]


def load_pdf(path: str) -> list[dict]:
    """Load a PDF file using pypdf directly."""
    reader = pypdf.PdfReader(path)
    docs = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append({"content": text, "metadata": {"source": path, "page": i}})
    return docs


def load_txt(path: str) -> list[dict]:
    """Load a plain text file."""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return [{"content": text, "metadata": {"source": path}}]


def extract_shape_images(shape):
    """Extract embedded images from a PowerPoint shape or group shape."""
    imgs = []
    try:
        if hasattr(shape, "image"):
            imgs.append(shape.image)
        elif getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # Group shape
            for sub_shape in shape.shapes:
                imgs.extend(extract_shape_images(sub_shape))
    except Exception:
        pass
    return imgs


def load_pptx(path: str, session_id: Optional[str] = None) -> Tuple[List[Dict], List[Dict]]:
    """
    Extract structured slides, bullet points, tables, and images from PowerPoint (.pptx).
    """
    # Deduce session_id from path if not explicitly passed
    if not session_id:
        norm_path = os.path.normpath(path)
        parts = norm_path.split(os.sep)
        if "uploaded_docs" in parts:
            idx = parts.index("uploaded_docs")
            if len(parts) > idx + 1:
                session_id = parts[idx + 1]

    prs = Presentation(path)
    docs = []
    slides_data = []
    for slide_idx, slide in enumerate(prs.slides):
        slide_texts = []
        title = ""
        bullets = []
        tables = []
        images = []

        # Check title shape first
        try:
            if slide.shapes.title and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
                slide_texts.append(title)
        except Exception:
            pass

        for shape in slide.shapes:
            try:
                if shape == getattr(slide.shapes, "title", None):
                    continue

                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)
                            if not title and len(text) < 80:
                                title = text
                            elif text != title:
                                bullets.append({
                                    "text": text,
                                    "level": getattr(paragraph, "level", 0),
                                })

                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        if any(row_cells):
                            table_rows.append(row_cells)
                            row_text = " | ".join(c for c in row_cells if c)
                            slide_texts.append(row_text)
                    if table_rows:
                        tables.append(table_rows)

                # Extract pictures/images
                shape_imgs = extract_shape_images(shape)
                for img in shape_imgs:
                    if len(images) >= 2:  # Cap at 2 images per slide (reduce RAM)
                        break
                    try:
                        pil_img = Image.open(io.BytesIO(img.blob))
                        try:
                            if pil_img.width > 500:
                                ratio = 500 / pil_img.width
                                new_size = (500, int(pil_img.height * ratio))
                                pil_img = pil_img.resize(new_size, Image.Resampling.LANCZOS)
                            pil_format = "PNG" if pil_img.mode in ("RGBA", "P") else "JPEG"

                            if session_id:
                                img_dir = os.path.join(UPLOAD_DIR, session_id, "slide_images")
                                os.makedirs(img_dir, exist_ok=True)
                                ext_name = "png" if pil_format == "PNG" else "jpg"
                                img_filename = f"slide_{slide_idx + 1}_img_{len(images) + 1}.{ext_name}"
                                img_file_path = os.path.join(img_dir, img_filename)
                                pil_img.save(img_file_path, format=pil_format, quality=85)
                                images.append(f"/slide_image/{session_id}/{img_filename}")
                            else:
                                buf = io.BytesIO()
                                pil_img.save(buf, format=pil_format, quality=85)
                                b64_str = base64.b64encode(buf.getvalue()).decode("utf-8")
                                mime = f"image/{pil_format.lower()}"
                                images.append(f"data:{mime};base64,{b64_str}")
                        finally:
                            pil_img.close()
                    except Exception as img_err:
                        print(f"Notice: Image extraction ({img_err})")
            except Exception:
                continue

        # Speaker notes
        notes = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass

        if not title and bullets:
            title = bullets.pop(0)
        if not title:
            title = f"Slide {slide_idx + 1}"

        slide_text = "\n".join(slide_texts)
        if notes:
            slide_text += f"\n\nSpeaker Notes: {notes}"

        docs.append({"content": slide_text, "metadata": {"source": path, "slide": slide_idx + 1}})

        slides_data.append({
            "slide_number": slide_idx + 1,
            "title": title if isinstance(title, str) else title.get("text", f"Slide {slide_idx + 1}"),
            "bullets": bullets,
            "tables": tables,
            "images": images,
            "notes": notes,
            "raw_text": slide_text,
        })

    prs = None
    gc.collect()
    return docs, slides_data


def load_file(path: str, ext: str, session_id: Optional[str] = None) -> Tuple[List[Dict], Optional[List[Dict]]]:
    """Dispatch file loading according to file extension."""
    if ext == ".pdf":
        return load_pdf(path), None
    elif ext == ".docx":
        return load_docx(path), None
    elif ext in (".pptx", ".ppt"):
        return load_pptx(path, session_id=session_id)
    elif ext == ".txt":
        return load_txt(path), None
    else:
        raise ValueError(f"Unsupported file extension: {ext}")
